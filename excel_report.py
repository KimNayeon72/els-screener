# -*- coding: utf-8 -*-
"""
그 주에 조회된 ELS 상품 중 수익률 15% 이상인 것만(필터는 main.py에서 미리
적용됨) 엑셀로 정리하고, 배리어 조건(2년간 하락 이력 없음)까지 만족하는
행에 색을 칠한다.
"""

import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial")
RECOMMEND_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RECOMMEND_FONT = Font(name="Arial", bold=True, color="006100")

COLUMNS = [
    "판매사", "종목명", "기초자산1", "기초자산2",
    "수익률(세전,연환산 %)", "최대손실률(%)",
    "1차조기상환배리어(%)", "배리어전체구조(1차~KI)", "조기상환주기/만기구조",
    "배리어확인방법",
    "만기", "만기일", "발행일(추정)", "청약마감일",
    "배리어2년내하락이력없음", "판정가능여부", "최종추천", "공시URL",
]


def build_report(rows: list, output_path: str = None) -> str:
    """rows: 각 상품에 스크리닝 결과가 합쳐진 dict 리스트.
    krx_client._parse_row() + enrich_with_kiwoom_barrier()가 만드는 스키마와
    underlying_price_checker가 추가하는 '배리어이하하락이력있음'(True/False/None)
    / '판정가능' 키를 기대한다.
    """
    if output_path is None:
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        output_path = os.path.join(OUTPUT_DIR, f"ELS_스크리닝_{date.today().isoformat()}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "ELS 스크리닝 결과"

    for col_idx, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r, item in enumerate(rows, start=2):
        underlyings = item.get("기초자산", [])
        u1 = underlyings[0] if len(underlyings) > 0 else ""
        u2 = underlyings[1] if len(underlyings) > 1 else ""

        yield_pct = item.get("수익률(세전, 연환산)")
        judgeable = bool(item.get("판정가능", False))

        breach = item.get("배리어이하하락이력있음")  # True / False / None(미확인)
        if breach is None:
            barrier_display = "확인불가"
            barrier_ok = False
        else:
            barrier_ok = not breach
            barrier_display = "O" if barrier_ok else "X"

        final_recommend = bool(judgeable and barrier_ok)
        applied_barrier = item.get("첫조기상환배리어(%)")
        applied_barrier_display = applied_barrier if applied_barrier is not None else "확인불가"
        barrier_structure = item.get("배리어전체구조") or "확인불가"
        period_structure = item.get("기간구조") or item.get("만기", "")

        values = [
            item.get("판매사", ""),
            item.get("종목명", ""),
            u1,
            u2,
            yield_pct if yield_pct is not None else "확인불가",
            item.get("최대손실률(%)", ""),
            applied_barrier_display,
            barrier_structure,
            period_structure,
            item.get("배리어출처", ""),
            item.get("만기", ""),
            item.get("만기일", ""),
            item.get("발행일(추정)", ""),
            item.get("청약마감일", ""),
            barrier_display,
            "O" if judgeable else "X",
            "★추천" if final_recommend else "",
            item.get("공시URL", ""),
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col_idx, value=val)
            cell.font = RECOMMEND_FONT if final_recommend else BODY_FONT
            if final_recommend:
                cell.fill = RECOMMEND_FILL

    for col_idx, header in enumerate(COLUMNS, start=1):
        max_len = max([len(str(header))] + [len(str(ws.cell(row=r, column=col_idx).value or ""))
                                             for r in range(2, ws.max_row + 1)])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 45)

    ws.freeze_panes = "A2"
    wb.save(output_path)
    return output_path
